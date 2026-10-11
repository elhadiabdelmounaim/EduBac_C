from datetime import timedelta
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from classrooms.models import Classroom, ClassroomMember
from classrooms.services.excel_results import class_excel_path, regenerate_class_excel
from education.models import Course, Lesson
from whiteboard.models import QuizCorrectionBoard
from .models import Quiz, Question, Choice, QuizAssignment, Attempt, StudentAnswer


@override_settings(ALLOWED_HOSTS=['testserver'])
class QuizDeletionTests(TestCase):
    def setUp(self):
        media = TemporaryDirectory()
        self.addCleanup(media.cleanup)
        override = override_settings(MEDIA_ROOT=media.name)
        override.enable()
        self.addCleanup(override.disable)
        self.teacher = User.objects.create_user(
            username='delete-teacher', email='delete-teacher@example.test', role='teacher',
        )
        self.other_teacher = User.objects.create_user(
            username='other-delete-teacher', email='other-delete-teacher@example.test', role='teacher',
        )
        self.student = User.objects.create_user(
            username='delete-student', email='delete-student@example.test', role='student',
        )
        self.classroom = Classroom.objects.create(name='Deletion test', teacher=self.teacher)
        ClassroomMember.objects.create(classroom=self.classroom, user=self.student)
        course = Course.objects.create(name='Test', niveau='tronc_commun')
        self.lesson = Lesson.objects.create(course=course, title='Leçon conservée')
        self.quiz = Quiz.objects.create(
            title='Quiz à supprimer', lesson=self.lesson, created_by=self.teacher,
            classroom=self.classroom, status='published',
        )
        self.question = Question.objects.create(quiz=self.quiz, text='1+1', correct_answer='2')
        self.choice = Choice.objects.create(question=self.question, text='2', is_correct=True)
        self.assignment = QuizAssignment.objects.create(quiz=self.quiz, student=self.student)
        self.attempt = Attempt.objects.create(
            quiz=self.quiz, student=self.student, submitted_at=timezone.now(), score=100,
        )
        self.answer = StudentAnswer.objects.create(
            attempt=self.attempt, question=self.question, answer='2', is_correct=True,
        )
        self.correction = QuizCorrectionBoard.objects.create(
            attempt=self.attempt, question=self.question,
        )
        self.keep = Quiz.objects.create(
            title='Autre quiz conservé', lesson=self.lesson, created_by=self.teacher,
        )
        self.url = reverse('quizzes:delete', args=[self.quiz.pk])
        self.login(self.teacher)

    def login(self, user, client=None):
        session = (client or self.client).session
        session['user_id'] = user.pk
        session.save()

    def test_confirmation_displays_real_counts_without_deleting(self):
        response = self.client.get(self.url)
        self.assertContains(response, 'Cette action est définitive')
        self.assertEqual(response.context['question_count'], 1)
        self.assertEqual(response.context['assignment_count'], 1)
        self.assertEqual(response.context['attempt_count'], 1)
        self.assertEqual(response.context['answer_count'], 1)
        self.assertTrue(Quiz.objects.filter(pk=self.quiz.pk).exists())
        self.assertContains(response, reverse('quizzes:manage'))

    def test_confirmed_delete_removes_related_data_but_not_accounts_classes_lessons_or_other_quizzes(self):
        response = self.client.post(self.url, {'confirm_delete': '1'}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Quiz.objects.filter(pk=self.quiz.pk).exists())
        for model, obj in (
            (Question, self.question), (Choice, self.choice),
            (QuizAssignment, self.assignment), (Attempt, self.attempt),
            (StudentAnswer, self.answer), (QuizCorrectionBoard, self.correction),
        ):
            self.assertFalse(model.objects.filter(pk=obj.pk).exists())
        self.assertTrue(User.objects.filter(pk=self.student.pk).exists())
        self.assertTrue(Classroom.objects.filter(pk=self.classroom.pk).exists())
        self.assertTrue(Lesson.objects.filter(pk=self.lesson.pk).exists())
        self.assertTrue(Quiz.objects.filter(pk=self.keep.pk).exists())

    def test_post_without_confirmation_does_not_delete(self):
        for data in ({}, {'confirm_delete': '0'}):
            response = self.client.post(self.url, data)
            self.assertContains(response, 'Confirmez la suppression', status_code=400)
        self.assertTrue(Quiz.objects.filter(pk=self.quiz.pk).exists())

    def test_old_quizzes_without_author_and_all_statuses_can_be_deleted(self):
        for status in ('draft', 'published', 'archived'):
            old = Quiz.objects.create(title=f'Ancien {status}', lesson=self.lesson, status=status)
            Quiz.objects.filter(pk=old.pk).update(created_at=timezone.now() - timedelta(days=365))
            url = reverse('quizzes:delete', args=[old.pk])
            self.assertEqual(self.client.get(url).status_code, 200)
            self.assertEqual(self.client.post(url, {'confirm_delete': '1'}).status_code, 302)
            self.assertFalse(Quiz.objects.filter(pk=old.pk).exists())

    def test_other_teacher_cannot_view_or_delete_owned_quiz(self):
        self.login(self.other_teacher)
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.assertEqual(self.client.post(self.url, {'confirm_delete': '1'}).status_code, 404)
        self.assertTrue(Quiz.objects.filter(pk=self.quiz.pk).exists())

    def test_student_and_anonymous_cannot_delete(self):
        self.login(self.student)
        for method in ('get', 'post'):
            response = getattr(self.client, method)(self.url, {'confirm_delete': '1'})
            self.assertRedirects(response, reverse('education:home'), fetch_redirect_response=False)
        self.client.logout()
        response = self.client.post(self.url, {'confirm_delete': '1'})
        self.assertRedirects(response, reverse('accounts:login'), fetch_redirect_response=False)
        self.assertTrue(Quiz.objects.filter(pk=self.quiz.pk).exists())

    def test_csrf_protection_and_repeated_deletion(self):
        client = Client(enforce_csrf_checks=True)
        self.login(self.teacher, client)
        self.assertEqual(client.post(self.url, {'confirm_delete': '1'}).status_code, 403)
        self.assertTrue(Quiz.objects.filter(pk=self.quiz.pk).exists())
        self.client.post(self.url, {'confirm_delete': '1'})
        self.assertEqual(self.client.post(self.url, {'confirm_delete': '1'}).status_code, 404)
        self.assertTrue(Quiz.objects.filter(pk=self.keep.pk).exists())

    def test_delete_links_on_teacher_lists_and_question_editor(self):
        for route in ('quizzes:manage', 'quizzes:training'):
            self.assertContains(self.client.get(reverse(route)), self.url)
        self.assertContains(
            self.client.get(reverse('quizzes:edit_questions', args=[self.quiz.pk])), self.url,
        )

    def test_cached_excel_invalidated_and_regenerated_without_deleted_quiz(self):
        regenerate_class_excel(self.classroom)
        path = class_excel_path(self.classroom)
        self.assertTrue(path.exists())
        self.client.post(self.url, {'confirm_delete': '1'})
        self.assertFalse(path.exists())
        response = self.client.get(reverse('classrooms:download_excel', args=[self.classroom.pk]))
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        self.assertEqual(workbook.sheetnames, ['_meta'])
        workbook.close()
        response.close()

    @patch('classrooms.services.excel_results.delete_class_excel', return_value=False)
    def test_failed_excel_invalidation_is_reported_without_hiding_success(self, delete_excel):
        regenerate_class_excel(self.classroom)
        response = self.client.post(self.url, {'confirm_delete': '1'}, follow=True)
        self.assertFalse(Quiz.objects.filter(pk=self.quiz.pk).exists())
        self.assertContains(response, 'certains exports Excel doivent être régénérés')
        delete_excel.assert_called_once()
